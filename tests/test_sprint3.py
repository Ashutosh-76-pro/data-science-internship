from __future__ import annotations

from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl import load_workbook

from src.screener.engine import add_composite_quality_score, apply_thresholds, load_config, preset
from src.analytics.peer import PEER_METRICS, compute_peer_percentiles, load_peer_groups

ROOT=Path(__file__).resolve().parents[1]
DF=pd.read_csv(ROOT/"data/n100_financials.csv")
PG=load_peer_groups(ROOT/"data/peer_groups.xlsx")
CONFIG=load_config(ROOT/"config/screener_config.yaml")


def test_01_universe_has_92_unique_companies():
    assert len(DF)==92 and DF.company_id.nunique()==92


def test_02_all_15_filterable_metrics_are_configured():
    assert len(CONFIG["custom_filters"]["supported_metrics"])==15


def test_03_all_six_presets_present():
    assert list(CONFIG["presets"]) == ["quality_compounder","value_pick","growth_accelerator","dividend_champion","debt_free_bluechip","turnaround_watch"]


def test_04_all_presets_return_between_5_and_50():
    counts={name:len(preset(DF,name,ROOT/"config/screener_config.yaml")) for name in CONFIG["presets"]}
    assert all(5<=n<=50 for n in counts.values()), counts


def test_05_quality_compounder_enforces_roe_and_de():
    out=preset(DF,"quality_compounder",ROOT/"config/screener_config.yaml")
    non_fin=out[out.broad_sector.ne("Financials")]
    assert (non_fin.roe_pct>15).all()
    assert (non_fin.debt_to_equity<1).all()


def test_06_financials_are_not_removed_by_de_filter():
    only_fin=DF[DF.broad_sector.eq("Financials")].copy()
    mask=apply_thresholds(only_fin,{"debt_to_equity":{"op":"lt","value":0.01}})
    assert mask.all()


def test_07_debt_free_icr_is_infinity():
    debt_free=DF[DF.debt_to_equity.eq(0)].iloc[0]
    from src.screener.engine import _numeric_series
    assert _numeric_series(pd.Series([debt_free.interest_coverage])).iloc[0] == float("inf")


def test_08_composite_score_is_0_to_100():
    score=add_composite_quality_score(DF)
    assert score.composite_quality_score.between(0,100).all()


def test_09_exactly_11_peer_groups_and_all_companies_assigned():
    assert PG.peer_group_name.nunique()==11
    assert PG.company_id.nunique()==92
    assert set(PG.company_id)==set(DF.company_id)


def test_10_ten_peer_metrics_are_defined():
    assert len(PEER_METRICS)==10


def test_11_peer_percentiles_has_920_records():
    score=add_composite_quality_score(DF)
    p=compute_peer_percentiles(score,PG,year=2024)
    assert len(p)==920
    assert p.percentile_rank.dropna().between(0,1).all()


def test_12_it_services_highest_roe_has_highest_percentile():
    score=add_composite_quality_score(DF)
    p=compute_peer_percentiles(score,PG,year=2024)
    it=p[(p.peer_group_name.eq("IT Services"))&(p.metric.eq("roe_pct"))].sort_values("percentile_rank",ascending=False)
    it_ids=PG[PG.peer_group_name.eq("IT Services")].company_id
    roe=score[score.company_id.isin(it_ids)].set_index("company_id").roe_pct
    assert it.iloc[0].company_id == roe.idxmax()


def test_13_generated_excel_sheet_counts_are_exact():
    wb1=load_workbook(ROOT/"output/screener_output.xlsx",read_only=True)
    wb2=load_workbook(ROOT/"output/peer_comparison.xlsx",read_only=True)
    assert len(wb1.sheetnames)==6
    assert len(wb2.sheetnames)==11


def test_14_sqlite_peer_percentiles_table_has_920_rows():
    with sqlite3.connect(ROOT/"output/screener.sqlite") as con:
        tables={r[0] for r in con.execute("select name from sqlite_master where type='table'")}
        n=con.execute("select count(*) from peer_percentiles").fetchone()[0]
    assert "peer_percentiles" in tables
    assert n==920
