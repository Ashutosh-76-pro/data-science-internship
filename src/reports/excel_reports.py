from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from src.screener.engine import COLUMN_LABELS, DISPLAY_KPI_COLUMNS

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GOLD = PatternFill("solid", fgColor="FFD966")
HEADER = PatternFill("solid", fgColor="1F4E78")
WHITE_FONT = Font(color="FFFFFF", bold=True)


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill=HEADER; cell.font=WHITE_FONT; cell.alignment=Alignment(horizontal="center")
    ws.auto_filter.ref=ws.dimensions
    for col in range(1, ws.max_column+1):
        max_len=0
        for row in ws.iter_rows(min_col=col,max_col=col,min_row=1,max_row=min(ws.max_row,100)):
            for c in row:
                max_len=max(max_len,len(str(c.value)) if c.value is not None else 0)
        ws.column_dimensions[get_column_letter(col)].width=min(max(max_len+2,10),28)


def export_screener(df: pd.DataFrame, preset_results: Mapping[str,pd.DataFrame], out_path: str | Path, config: dict):
    out=Path(out_path); out.parent.mkdir(parents=True,exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for preset_name, result in preset_results.items():
            cols=["company_id","company_name","peer_group_name","broad_sector","year"]+DISPLAY_KPI_COLUMNS+["composite_quality_score"]
            frame=result[[c for c in cols if c in result.columns]].copy()
            frame=frame.rename(columns={c:COLUMN_LABELS.get(c,c) for c in frame.columns})
            sheet=preset_name.replace("_"," ").title()[:31]
            frame.to_excel(writer,sheet_name=sheet,index=False)
    wb=load_workbook(out)
    for preset_name,result in preset_results.items():
        ws=wb[preset_name.replace("_"," ").title()[:31]]
        thresholds=config["presets"][preset_name]
        header={str(c.value):c.column for c in ws[1]}
        for row_idx, (_, row) in enumerate(result.iterrows(), start=2):
            for metric, rule in thresholds.items():
                label=COLUMN_LABELS.get(metric,metric)
                col=header.get(label)
                if not col: continue
                if metric=="debt_to_equity" and str(row.get("broad_sector","")).casefold()=="financials":
                    continue
                try:
                    val=row[metric]; target=float(rule["value"]); op=rule["op"]
                    if op in ("gt", ">"): passed=float(val)>target
                    elif op in ("lt", "<"): passed=float(val)<target
                    elif op in ("eq","="): passed=abs(float(val)-target)<1e-12
                    elif op in ("min",">="): passed=float(val)>=target
                    else: passed=float(val)<=target
                except Exception:
                    passed=False
                ws.cell(row_idx,col).fill=GREEN if passed else RED
        _style_sheet(ws)
    wb.save(out)


def export_peer(df: pd.DataFrame, percentiles: pd.DataFrame, peer_groups: pd.DataFrame, out_path: str | Path):
    out=Path(out_path); out.parent.mkdir(parents=True,exist_ok=True)
    metric_order=list(__import__("src.analytics.peer",fromlist=["PEER_METRICS"]).PEER_METRICS)
    value_cols=["company_id","company_name","broad_sector","year"]+DISPLAY_KPI_COLUMNS
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for group in peer_groups["peer_group_name"].drop_duplicates().tolist():
            g_ids=peer_groups.loc[peer_groups.peer_group_name.eq(group),"company_id"]
            frame=df[df.company_id.isin(g_ids)].copy()
            piv=percentiles[percentiles.peer_group_name.eq(group)].pivot(index="company_id",columns="metric",values="percentile_rank")
            piv=piv.reindex(columns=metric_order)
            piv.columns=[f"{COLUMN_LABELS.get(c,c)} Percentile" for c in piv.columns]
            frame=frame.set_index("company_id").join(piv,how="left").reset_index()
            cols=[c for c in value_cols if c in frame.columns]+list(piv.columns)
            frame=frame[cols]
            frame.to_excel(writer,sheet_name=group[:31],index=False)
    wb=load_workbook(out)
    for group in peer_groups["peer_group_name"].drop_duplicates().tolist():
        ws=wb[group[:31]]; _style_sheet(ws)
        header={str(c.value):c.column for c in ws[1]}
        rank_cols=[c for c in ws[1] if str(c.value).endswith(" Percentile")]
        for row in range(2,ws.max_row+1):
            company=ws.cell(row,header["company_id"]).value
            is_benchmark=bool(peer_groups.loc[(peer_groups.peer_group_name.eq(group))&(peer_groups.company_id.eq(company)),"benchmark_company"].iloc[0])
            for c in rank_cols:
                val=ws.cell(row,c.column).value
                if val is None: continue
                if float(val)>=0.75: ws.cell(row,c.column).fill=GREEN
                elif float(val)<=0.25: ws.cell(row,c.column).fill=RED
                else: ws.cell(row,c.column).fill=YELLOW
            if is_benchmark:
                for col in range(1,ws.max_column+1): ws.cell(row,col).fill=GOLD
        summary_row=ws.max_row+2
        ws.cell(summary_row,1,"Peer Group Median")
        for col in range(2,ws.max_column+1):
            vals=[ws.cell(row,col).value for row in range(2,ws.max_row+1) if isinstance(ws.cell(row,col).value,(int,float)) and not isinstance(ws.cell(row,col).value,bool)]
            if vals:
                import statistics
                ws.cell(summary_row,col,statistics.median(vals))
        for c in ws[summary_row]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="D9EAF7")
        ws.freeze_panes="A2"
        for col in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width=min(max(ws.column_dimensions[get_column_letter(col)].width or 10,12),28)
    wb.save(out)
